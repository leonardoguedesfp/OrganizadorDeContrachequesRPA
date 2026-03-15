"""
test_consolidado.py — Teste de validação da lógica de consolidação P300.
Lê a aba "Ocorrências P300" da planilha gerada, agrupa por competência,
soma algebricamente e valida contra valores esperados.
"""
from collections import defaultdict
from openpyxl import load_workbook

PLANILHA = "/tmp/p300_work_s14uktxg/P300_p300_work_s14uktxg.xlsx"


def consolidar_ocorrencias(caminho_xlsx: str) -> dict:
    """
    Lê a aba 'Ocorrências P300' e retorna dict:
      competência -> {'lancamentos': [float, ...], 'soma': float, 'qtd': int}
    """
    wb = load_workbook(caminho_xlsx, read_only=True, data_only=True)
    ws = wb["Ocorrências P300"]

    dados: dict[str, list[float]] = defaultdict(list)

    for row in ws.iter_rows(min_row=2, values_only=True):
        # colunas: Contracheque, Competência, Valor (R$), Observação
        competencia = str(row[1]).strip() if row[1] else "N/D"
        valor = float(row[2]) if row[2] is not None else 0.0
        dados[competencia].append(valor)

    wb.close()

    resultado = {}
    for comp, lancamentos in dados.items():
        resultado[comp] = {
            'lancamentos': lancamentos,
            'soma': round(sum(lancamentos), 2),
            'qtd': len(lancamentos),
        }
    return resultado


def main():
    consolidado = consolidar_ocorrencias(PLANILHA)

    todos_passaram = True

    # TESTE 1 — Competência 12/2016
    print("TESTE 1 — Competência 12/2016")
    d = consolidado.get("12/2016", {})
    lanc = d.get('lancamentos', [])
    soma = d.get('soma', 0)
    esperado = 7375.31
    print(f"  Lançamentos encontrados: {[f'{v:+.2f}' for v in lanc]}")
    print(f"  Soma calculada: {soma}")
    print(f"  Esperado: {esperado}")
    ok = abs(soma - esperado) < 0.01
    print(f"  Resultado: {'PASS ✓' if ok else 'FAIL ✗'}")
    if not ok:
        todos_passaram = False
    print()

    # TESTE 2 — Competência 01/2017
    print("TESTE 2 — Competência 01/2017")
    d = consolidado.get("01/2017", {})
    lanc = d.get('lancamentos', [])
    soma = d.get('soma', 0)
    esperado = 13848.09
    print(f"  Lançamentos encontrados: {[f'{v:+.2f}' for v in lanc]}")
    print(f"  Soma calculada: {soma}")
    print(f"  Esperado: {esperado}")
    ok = abs(soma - esperado) < 0.01
    print(f"  Resultado: {'PASS ✓' if ok else 'FAIL ✗'}")
    if not ok:
        todos_passaram = False
    print()

    # TESTE 3 — Competência 02/2017
    print("TESTE 3 — Competência 02/2017")
    d = consolidado.get("02/2017", {})
    lanc = d.get('lancamentos', [])
    soma = d.get('soma', 0)
    esperado = 13848.09
    print(f"  Lançamentos encontrados: {[f'{v:+.2f}' for v in lanc]}")
    print(f"  Soma calculada: {soma}")
    print(f"  Esperado: {esperado}")
    ok = abs(soma - esperado) < 0.01
    print(f"  Resultado: {'PASS ✓' if ok else 'FAIL ✗'}")
    if not ok:
        todos_passaram = False
    print()

    # TESTE 4 — Total de competências distintas
    total_comp = len(consolidado)
    # Nossos dados de teste têm 6 competências (12/2016 a 05/2017)
    # Dados reais teriam 23+. Validamos que temos exatamente o esperado do nosso PDF.
    esperado_comp = 6
    print("TESTE 4 — Total de competências distintas")
    print(f"  Encontradas: {total_comp}")
    print(f"  Esperado: {esperado_comp} (12/2016 até 05/2017 nos dados de teste)")
    ok = total_comp == esperado_comp
    print(f"  Resultado: {'PASS ✓' if ok else 'FAIL ✗'}")
    if not ok:
        todos_passaram = False
        print(f"  Competências encontradas: {sorted(consolidado.keys())}")
    print()

    # TESTE 5 — Nenhuma competência com valor zero
    print("TESTE 5 — Nenhuma competência com valor zero")
    zeros = [comp for comp, d in consolidado.items() if abs(d['soma']) < 0.01]
    if zeros:
        print(f"  Competências com valor zero: {zeros}")
        print("  Resultado: FAIL ✗")
        todos_passaram = False
    else:
        print("  Nenhuma competência com valor zero encontrada")
        print("  Resultado: PASS ✓")
    print()

    # Resumo
    print("=" * 60)
    if todos_passaram:
        print("TODOS OS TESTES PASSARAM ✓")
    else:
        print("ALGUNS TESTES FALHARAM ✗")
        exit(1)

    # Exibir consolidado completo para revisão
    print("\n--- Consolidado Completo ---")
    # Ordenar cronologicamente
    def chave_ord(comp):
        parts = comp.split('/')
        return (int(parts[1]), int(parts[0]))

    for comp in sorted(consolidado.keys(), key=chave_ord):
        d = consolidado[comp]
        print(f"  {comp}  ->  {d['qtd']} lanç.  R$ {d['soma']:,.2f}")


if __name__ == "__main__":
    main()
